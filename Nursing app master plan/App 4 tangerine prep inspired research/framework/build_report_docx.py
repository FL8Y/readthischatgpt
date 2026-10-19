import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

A2 = '/tmp/claude-0/-home-user/386c8b50-43be-54c5-b6a0-357edf2ecf6f/scratchpad/tan/'
R = '/home/user/readthischatgpt/Nursing app master plan/App 4 tangerine prep inspired research/'

d = Document()
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.9)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); tcPr.append(sh)

def H(t, l=1): d.add_heading(t, level=l)
def P(t, bold=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic; return p
def B(items):
    for t in items:
        p = d.add_paragraph(style='List Bullet')
        if isinstance(t, tuple):
            r = p.add_run(t[0]); r.bold = True; p.add_run(t[1])
        else: p.add_run(t)
def T(head, rows, widths=None):
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); shade(c, 'D9E2F3')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    d.add_paragraph()
def CODE(t):
    p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.15)
    r = p.add_run(t); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    pPr = p._p.get_or_add_pPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'F2F2F2'); pPr.append(sh)
def NOTE(t):
    p = d.add_paragraph(); r = p.add_run('Note: '); r.bold = True; p.add_run(t)


def inline(p, text):
    # **bold** and `code` runs
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else: p.add_run(part)

def MD(path, level_shift=1):
    lines = open(path, errors='replace').read().split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln or ln.startswith('Labels:') or ln.startswith('Source is untrusted'): i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            lvl = min(len(m.group(1)) + level_shift, 4)
            if len(m.group(1)) == 1: i += 1; continue   # file title replaced by our own section heading
            d.add_heading(m.group(2), level=lvl); i += 1; continue
        if ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells): rows.append(cells)
                i += 1
            if rows:
                n = max(len(r) for r in rows)
                t = d.add_table(rows=0, cols=n); t.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(n):
                        cells[ci].text = ''
                        p = cells[ci].paragraphs[0]
                        inline(p, r[ci] if ci < len(r) else '')
                        for run in p.runs: run.font.size = Pt(9); run.bold = run.bold or ri == 0
                        if ri == 0: shade(cells[ci], 'D9E2F3')
                d.add_paragraph()
            continue
        mb = re.match(r'^(\s*)[-*]\s+(.*)', ln)
        if mb:
            p = d.add_paragraph(style='List Bullet'); inline(p, mb.group(2)); i += 1; continue
        p = d.add_paragraph(); inline(p, ln.strip()); i += 1


def FILE(path, maxlines=None):
    t = open(path, encoding='utf-8').read().rstrip('\n')
    CODE(t)


# ---------- Title ----------
d.add_heading('Tangerine Prep — How the App Works', 0)
P('Static reverse-engineering report of Tangerine Prep 1.9.0 (Android package com.tangerineprep.app), a nursing licensure-exam prep app, with lessons for the Nursing app master plan and a tested reference engine.', italic=True)
P('Source: Tanrinee_part1_of_3 … part3_of_3, three byte-split pieces of one file. Joined in order they form a 71 MB 7-Zip archive containing base.apk (123 MB). Static analysis only: the app was never run and no requests were made to its servers. Fonts were excluded.')

H('1. Summary')
B([
 ('What it is. ', 'A mobile exam-prep app for the Philippine Nurse Licensure Examination (PNLE; NCLEX also mentioned), with AI-curated practice sets, mock board exams, an adaptive exam, NCLEX-style case studies, flashcards, a study planner, an AI tutor persona called Cleo, battles between friends, social groups, streaks and a paid tier.'),
 ('How it is built. ', 'An Expo SDK 54 / React Native app. Almost all logic is a Hermes v96 bytecode bundle (9.6 MB). This version was fully decompiled to readable pseudo-JavaScript (56 MB), so unlike the Gizmo report, real code paths could be followed.'),
 ('Where the brains are. ', 'Question scoring, adaptive selection, spaced repetition, ranking and AI generation are all on the server (a Django-style REST API on Railway). The app is a client that sends answers and renders results; the formulas are not in the APK.'),
 ('What to take from it. ', 'Section 9 turns the findings into a feature blueprint for your nursing app: content taxonomy, item types, practice modes, energy and streak mechanics, offline mode and pricing structure. Section 10 delivers a tested TypeScript engine for scoring, adaptive selection, streaks and a free-tier cap.'),
 ('Privacy finding. ', 'The bundle ships about 28,652 full names of exam passers (for a results-lookup feature) and uses Microsoft Clarity session replay plus Meta and Google ad attribution. The names are deliberately not copied into this report or the repo. If you build something similar, handle licensure-result data with consent and care.'),
])

H('2. What was uploaded')
T(['Item', 'Value'], [
 ['Parts', 'Tanrinee_part1_of_3 … part3_of_3 (23.8 MB each; each is a zip holding one .bin piece)'],
 ['Joined', 'cat part1 part2 part3 = 71,303,052 bytes, a 7-Zip archive'],
 ['Content', 'base.apk, 123 MB, 1,434 entries, 7 dex files, one Hermes bundle, no lib/ folder in the upload'],
 ['Package / version', 'com.tangerineprep.app, 1.9.0 (versionCode 19), minSdk 24, targetSdk 36'],
 ['Fonts', 'Excluded as requested'],
])

H('3. How the analysis was done and its limits')
B([
 'Rebuilt the file from its parts, extracted the APK, decoded the manifest with apktool, read app.config and the Expo manifest.',
 'The Hermes bundle is bytecode version 96, which the hermes-dec decompiler supports, so the whole bundle was decompiled to pseudo-JS and its 88,668 strings were dumped. Variable names are partly lost, but routes, API paths, constants and logic are readable.',
 'Three parallel research passes (Sonnet agents) covered the backend and network layer, the learning features, and the native/SDK/security side. Their notes are Addenda A–C and also saved as Markdown in the folder.',
])
NOTE('"Observed" is a literal string, constant or field in the code. "Inferred" is my reading of how the pieces fit. Scoring, adaptive and SRS algorithms are server-side, so their exact formulas are unknown. Native .so libraries were not in the upload.')

H('4. App identity and stack')
T(['Item', 'Value', 'Type'], [
 ['Name / slug', 'Tangerine Prep / tangerine-prep', 'observed'],
 ['Package / scheme', 'com.tangerineprep.app / tangerineprep://', 'observed'],
 ['Framework', 'Expo SDK 54, React Native new architecture, expo-router, Hermes bytecode v96', 'observed'],
 ['Backend', 'https://tangerine-backend-production.up.railway.app/api (primary); https://tangerineprep.app/api (fallback); Django-REST-style trailing slashes', 'observed / inferred'],
 ['Auth', 'Email, Apple and Google sign-in; JWT access and refresh tokens kept in SecureStore', 'observed'],
 ['Realtime', 'WebSockets (4 endpoints) and LiveKit/WebRTC for study and battle voice', 'observed'],
 ['Payments', 'Play Billing 8.3 via react-native-iap / OpenIAP; monthly, quarterly, yearly; one-time QR Ph passes on the web', 'observed'],
 ['Notifications', 'Expo push + FCM, Notifee for local scheduling', 'observed'],
 ['Analytics', 'Microsoft Clarity (session replay), Facebook SDK, first-party /app-platform/events/ and /errors/', 'observed'],
 ['Anti-tamper', 'Google Play pairip license check (application class com.pairip.application.Application)', 'observed'],
])

H('5. Architecture')
CODE("""+----------------------------------------------------------+
|  Android app (Expo / React Native, Hermes bundle)        |
|  expo-router screens - local state - SQLite              |
|  offline bundle + outbox - Notifee local reminders       |
+---------------+------------------------+-----------------+
                | HTTPS (Bearer JWT)     | WebSocket / LiveKit
                v                        v
+---------------+------------------------+-----------------+
|  Railway backend (Django REST style, /api/...)           |
|  auth - billing - questions - sessions - stats           |
|  study-plan - flashcards (server SRS) - social - battles |
|  AI generation (questions, flashcards, Cleo tutor)       |
+----------------------------------------------------------+
   Failover: request loop tries the next base URL on network
   error or HTTP 502/503/504.""")
P('Design points worth copying (observed): base-URL failover, a token-refresh call, a chunked offline bundle with a local outbox that syncs when the connection returns, and first-party error and event logging.')

H('6. How the learning features work')
T(['Feature', 'What it does (observed unless marked)'], [
 ['Content taxonomy', 'PNLE content is organised NP1–NP5, each split into practice area > topic > TOS (table of specifications) section. Endpoints /practice-areas/ and /tos/question-structure/ supply the tree.'],
 ['Practice modes', 'Custom (choose areas and count), AI (Cleo curates a set from the existing question bank; it does not write new exam questions), Mock Board, Adaptive Exam (a learner model called "Chev"), and a coaching waitlist pilot.'],
 ['Case studies', 'Pre-authored cases served from /situations/v2/. Their items are mcq, sata (select all that apply), matrix, highlight, cloze and free_response, scored in credits with partial credit. Ordinary practice questions are single-select only.'],
 ['Rationales', 'Every question carries a rationale; there is a rationale-video lab screen for video explanations.'],
 ['Flashcards', 'Folders, packs, a card editor and shared decks. Review uses four buttons (again / hard / good / easy); the SRS is server-side. A review-undo endpoint exists.'],
 ['Study planner', '/study-plan/v2/ status, calendar, events and rebalance: it builds a schedule to the exam date and can rebalance it when you fall behind (algorithm server-side).'],
 ['Cheat sheets / library', 'Condensed PNLE notes; a library of your own materials with folders, sources, search and AI refine.'],
 ['AI tutor "Cleo"', 'Ask-Cleo with streaming answers, voice transcription, usage limits and visuals (endpoints under /sessions/ask-cleo/v2/).'],
 ['Diagnostics and reports', 'Diagnostic test, passing-estimate and passing-probability with history and rank neighbours, a Bloom-level exam report with a premium tier.'],
 ['Social', 'Friends, groups, direct messages, study rooms with voice (the deeper pass found no leaderboard list; see the Complete Guide), a "townhall" community, and a PNLE results lookup/celebration feature.'],
 ['Battles', 'Duel and Royale modes with speed and streak bonuses, a ladder from Student II up to Apex Topnotcher, and 90-day seasons.'],
 ['Gamification', 'Streak days on a Manila-time day clock, a daily goal, and "energy" (free tier capped at 20). There is no XP counter and no spendable currency despite the tangerine branding.'],
 ['Offline mode', 'Chunked SQLite bundle (/offline/bundle/manifest and chunks) with an outbox table for later sync.'],
])

H('7. Pricing and monetisation')
B([
 'Free tier with a daily energy cap of 20 questions; premium removes it.',
 'Subscriptions in pesos: observed price points of 399, 799 and 1,999 for the monthly, quarterly and yearly plans (the pairing in that order is inferred), plus one-time QR Ph passes for local payment.',
 'Play Billing purchases are verified server-side through /auth/billing/native-iap/sync/ and /v1/subscriptions/*.',
 'A premium exam-analysis report exists (/exam-report/analysis/premium); AI usage appears to be metered (an ask-cleo usage endpoint exists).',
])

H('8. Privacy, security and compliance observations')
B([
 ('Exam-passer names. ', 'The bundle contains roughly 28,652 full names for the August 2026 PNLE results lookup and a "celebrate" consent flow. Public lists of licensure results exist, but embedding them in every install is notable. Names are not reproduced here.'),
 ('Session replay and ad tracking. ', 'Microsoft Clarity records sessions and the Facebook SDK collects advertising IDs (with Privacy Sandbox ad permissions). Chat and profile content could reach third parties unless masked. Students\' study data and chats raise consent questions.'),
 ('Backup and builds. ', 'allowBackup is true (only SecureStore excluded). The release build still contains expo-dev-client classes and an exported Compose PreviewActivity, which are developer leftovers.'),
 ('Anti-tamper. ', 'pairip license verification is present, so repackaging the app would fail the Play license check.'),
 ('Screen capture. ', 'DETECT_SCREEN_CAPTURE and expo-screen-capture are used to detect screenshots, probably to protect question content.'),
 ('Permissions. ', 'Microphone, camera, Bluetooth, notifications, billing, ad ID, system alert window. No location or storage permissions.'),
])

H('9. What this means for the Nursing app master plan')
P('A feature blueprint, ordered by value for effort:', bold=True)
T(['Priority', 'Feature', 'Why / how (see code in section 10)'], [
 ['1', 'Exam-blueprint taxonomy', 'Organise every item by exam blueprint (area > topic > section). It drives practice, analytics, planning and cheat sheets.'],
 ['1', 'NCLEX-style item types with partial credit', 'mcq, SATA, matrix, cloze, highlight. SATA and matrix scoring are in the reference engine.'],
 ['1', 'Rationales for every item', 'The app treats rationale as core; make it mandatory in your content schema.'],
 ['2', 'Spaced-repetition flashcards', 'Use the FSRS scheduler from the App 2 research; four-button review with undo.'],
 ['2', 'Adaptive practice and readiness estimate', 'Rasch-style ability tracking and "readiness" are in the reference engine; replace with calibrated item difficulty from real data.'],
 ['2', 'Study planner with rebalance', 'Plan backwards from the exam date by weakness; rebalance on missed days.'],
 ['2', 'Streak and daily goal on a fixed timezone', 'Implemented and tested; the fixed-timezone day clock avoids unfair streak breaks.'],
 ['3', 'AI question and flashcard generation', 'Pair with the AI next-quiz pipeline from the App 3 research; require clinician review before publishing.'],
 ['3', 'Offline bundle and outbox', 'Chunked content download, local SQLite, queued submissions.'],
 ['3', 'Social: groups, friends, leaderboards, battles', 'Highest effort, highest retention; add after the core loop works.'],
 ['4', 'Voice study rooms (LiveKit)', 'Optional; peer-to-peer voice token endpoint per room.'],
])
P('Pitfalls the app suggests avoiding:', bold=True)
B([
 'Do not ship answer keys or rationales in the client if you want to monetise content; keep scoring on the server (as Tangerine does) or accept that they can be extracted.',
 'Do not embed personal data lists in the app bundle.',
 'Be deliberate about session replay and ad SDKs in an app used by students; mask content and get consent.',
 'Remove developer tooling (dev client, preview activities) from release builds.',
 'Any AI-generated clinical content needs expert review and a visible disclaimer; exam-prep is not clinical guidance.',
])

H('10. Reference code: practice engine (TypeScript)')
P('Written for this project, informed by the features above. The app\'s own formulas are server-side and were not visible, so this is an independent design, not a copy. It provides: partial-credit scoring for mcq, SATA, matrix and cloze items; Rasch (1-parameter IRT) ability updates and next-item selection with a weak-area bonus; a readiness estimate; a streak counter on an IANA-timezone day clock (for example Asia/Manila); and a free-tier daily energy cap.')
P('engine.ts', bold=True)
FILE(R+'product/practice-engine/src/engine.ts')
P('Verification: 10 tests pass (node --test): scoring rules including the SATA +/- floor-at-zero rule, input validation, ability updates and bounds, adaptive pick, a simulated learner that converges toward its true level, timezone-dependent day keys, streak continue/reset/lapse, and the energy cap with daily reset. Not run in an app or against real learner data.')
NOTE('Calibrate before use: item difficulty must come from real response data, and the 20-per-day energy cap is a product decision you should tune.')

H('11. What could not be determined')
B([
 'Server-side algorithms: scoring weights, adaptive item selection ("Chev"), SRS parameters, planner rebalancing, ranking, AI prompts and models.',
 'Exact HTTP methods and response schemas for most of the ~140 endpoints (paths and several field names are observed).',
 'Content volume: question banks are fetched from the server; only a handful of demo questions are bundled.',
 'Native behaviour in LiveKit, WebRTC and Hermes (no lib/ folder in the upload).',
 'Whether Clarity masking is configured, and how the passer-name list is licensed.',
])

H('12. Addendum A — Backend, auth and voice')
MD(A2+'addendum_backend.md')
H('13. Addendum B — Learning features and logic')
MD(A2+'addendum_features.md')
H('14. Addendum C — Native, SDKs and security')
MD(A2+'addendum_native.md')
d.save(R+'Tangerine_Prep_Research.docx')
print('saved')
