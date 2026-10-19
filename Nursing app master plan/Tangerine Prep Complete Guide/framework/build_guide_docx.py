import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

A2 = '/home/user/readthischatgpt/Nursing app master plan/Tangerine Prep Complete Guide/research/'
R = '/home/user/readthischatgpt/Nursing app master plan/Tangerine Prep Complete Guide/'

d = Document()
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.9)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); tcPr.append(sh)

def H(t, l=1): d.add_heading(t, level=l)
def P(t, bold=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic; return p
def B(items):
    for t in items:
        p = d.add_paragraph(style='List Bullet')
        if isinstance(t, tuple):
            r = p.add_run(t[0]); r.bold = True; p.add_run(t[1])
        else: p.add_run(t)
def T(head, rows, widths=None):
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); shade(c, 'D9E2F3')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    d.add_paragraph()
def CODE(t):
    p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.15)
    r = p.add_run(t); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    pPr = p._p.get_or_add_pPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'F2F2F2'); pPr.append(sh)
def NOTE(t):
    p = d.add_paragraph(); r = p.add_run('Note: '); r.bold = True; p.add_run(t)


def inline(p, text):
    # **bold** and `code` runs
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else: p.add_run(part)

def MD(path, level_shift=1):
    lines = open(path, errors='replace').read().split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln or ln.startswith('Labels:') or ln.startswith('Source is untrusted'): i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            lvl = min(len(m.group(1)) + level_shift, 4)
            if len(m.group(1)) == 1: i += 1; continue   # file title replaced by our own section heading
            d.add_heading(m.group(2), level=lvl); i += 1; continue
        if ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells): rows.append(cells)
                i += 1
            if rows:
                n = max(len(r) for r in rows)
                t = d.add_table(rows=0, cols=n); t.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(n):
                        cells[ci].text = ''
                        p = cells[ci].paragraphs[0]
                        inline(p, r[ci] if ci < len(r) else '')
                        for run in p.runs: run.font.size = Pt(9); run.bold = run.bold or ri == 0
                        if ri == 0: shade(cells[ci], 'D9E2F3')
                d.add_paragraph()
            continue
        mb = re.match(r'^(\s*)[-*]\s+(.*)', ln)
        if mb:
            p = d.add_paragraph(style='List Bullet'); inline(p, mb.group(2)); i += 1; continue
        p = d.add_paragraph(); inline(p, ln.strip()); i += 1


def FILE(path, maxlines=None):
    t = open(path, encoding='utf-8').read().rstrip('\n')
    CODE(t)




def FILES(paths):
    for p in paths:
        P(p, bold=True); FILE(R+p)

d.add_heading('Tangerine Prep — Complete Guide', 0)
P('How every part of the app works: questions and practice, results and stats, history, diagnostic, learning path, cheat sheets, case studies, flashcards, the AI tutor, the library, pre-built content, buttons and navigation, social and battles, the UI and design system, client-side functions, pricing — plus a tested reference code kit.', italic=True)
P('Target: Tangerine Prep 1.9.0 (com.tangerineprep.app), a Philippine nursing licensure (PNLE) exam-prep app. Method: the app\'s JavaScript was fully decompiled (Hermes bytecode v96) and read screen by screen; seven research passes (Sonnet agents) each covered one feature group, with line references into the decompiled code. Static analysis only: the app was not run and no requests were made to its servers. No images or art were copied, and the list of exam-passer names embedded in the app is deliberately not reproduced.')

H('1. How to read this guide')
B([
 'Every claim is labelled OBSERVED (a literal string, constant or call in the code), INFERRED (my reading of how pieces fit) or SERVER (the logic runs on the server and is not visible). L-numbers are line references into the decompiled file (kept in the regenerable method, not in the repo).',
 'Parts A to G are the detailed feature guides, one per area. Part 6 is the reference code kit. Part 7 is the take-away for the nursing app.',
 'Where a later pass corrected an earlier note, Section 4 lists the correction. Trust the guides over the older App 4 report where they differ.',
])

H('2. The app in one picture')
CODE("""FIRST OPEN
  splash > landing > "Cleo" welcome (first name) > exam date > rest days > school > level choice
        |                                                            |
        |   optional DIAGNOSTIC (guest)  -> level label (40/60/80 %)  |
        v                                                            v
  study-plan reveal  ->  sign-up / sign-in  ->  notification permission (asked later, after login)

HOME (tab bar, 5 slots, Versus in the centre)
  Practice hub ----------- Custom | Mock Board | Adaptive | Topic preview | AI curator (Cleo) | Coaching waitlist
        |                          \\__ session player: answer, eliminate, pause, rationale, report problem, tag mastery
        v
  Results / report  -> exam analysis (premium) -> history -> "My stats" (areas, topics, passing estimate)
  Learning path (study plan) <- diagnostic + exam date; tasks launch sessions; rebalance when behind
  Study tools:  Cheat sheets | Case studies (situations) | Flashcards (server SRS) | Ask Cleo | Library
  Social:       friends | groups | DMs | study rooms (voice) | Townhall | Versus battles + rank ladder
  Money:        energy (free cap) | paywalls | monthly / quarterly / yearly | web checkout / passes""")

H('3. Feature map')
T(['Feature', 'Typical routes', 'Main API families', 'Logic lives', 'Detail'], [
 ['Practice setup and player', '/practice, /session/[id]', '/sessions/*, /questions/*', 'Selection, scoring, adaptive: server. Timer, UI: client', 'Part A'],
 ['Results, history, my stats', '/report/[id], /history, /performance', '/stats/*, /exam-report/*', 'Mostly server; report assembled client-side', 'Part A'],
 ['Onboarding and diagnostic', '/(auth)/welcome, /(auth)/diagnostic', '/diagnostics/*, /app-platform/onboarding-config/', 'Level bucket client-side; analysis server', 'Part B'],
 ['Learning path', '/study, /study/rebalance', '/study-plan/v2/*', 'Plan and rebalance: server', 'Part B'],
 ['Cheat sheets', '/cheat-sheets', '/study-sheets/, offline cache', 'Content server; reader client', 'Part B'],
 ['Case studies', '/case-studies, /situation-player', '/situations/v2/*', 'Grading server; five item types', 'Part B'],
 ['Flashcards', '/flashcards/*', '/flashcards/*', 'Server SRS; local queue helpers', 'Part B'],
 ['Ask Cleo', '/ask-cleo', '/sessions/ask-cleo/v2/*', 'LLM and curation: server', 'Part C'],
 ['Library', '/library/*', '/library/*', 'Generation: server; appears gated', 'Part C'],
 ['Pre-built content, paywall, analytics', 'various', '/offline/bundle/*, /subscription*', 'Content pipeline server-side', 'Part D'],
 ['Navigation, social, battles', 'tab bar, /townhall/*, /versus', 'WebSockets (4), /social/*', 'Matchmaking and points: server', 'Part E'],
 ['UI and design system', 'all', 'n/a', 'Client', 'Part F'],
 ['Client functions', 'all', 'n/a', 'Client', 'Part G'],
])

H('4. Corrections and open conflicts between passes')
P('The deeper passes refined or corrected earlier notes. Where they conflict, the later, more specific note wins.', italic=True)
T(['Topic', 'Earlier statement', 'Corrected finding'], [
 ['AI exam builder', '"AI generates questions"', 'Cleo curates existing bank questions; the library feature generates from user sources on the server'],
 ['Case-study item types', 'Six types including free_response', 'Five types in case studies; free_response belongs to class work'],
 ['Flag icon in the player', 'Report a question', 'Opens the mastery tagger; reporting uses a separate "Report a problem" sheet'],
 ['Report and review screens', 'Server report', 'The report is built client-side from the session detail; /review/[id] is a legacy redirect'],
 ['Time limit', 'Hard cut-off', 'Soft: at zero the app asks "Submit now" or "Continue answering"'],
 ['Settings and profile', '/settings and /profile routes', 'No such routes: settings is /townhall/profile-settings via a More menu; profile is the My profile tab in Townhall. No language, quiet-hours, delete-account or linked-accounts controls found'],
 ['Leaderboards', 'School and friends leaderboards', 'No leaderboard list found; there is a rank ladder, streaks and battles. No streak freeze and no rematch button found'],
 ['Flashcard review request', 'Sends state before and after', 'Sends card_id, rating and review_duration_ms only; state fields are telemetry'],
 ['Cleo "streaming"', 'Token streaming', 'The client reveals the reply paragraph by paragraph; the endpoint name says stream'],
 ['Offline integrity', 'Hashed chunks', 'Only version and index are checked against the manifest; no cryptographic hash found'],
 ['Library availability', 'Available feature', 'Appears gated to a TestFlight channel (inferred)'],
 ['Rank point maths', 'Maximum 39 / apex floor 40', 'Star counts give tier starts 0,4,8,12,16,20,25,30,35; the earlier maximum conflicts. The kit uses the star-count rule'],
 ['Notification permission', 'Asked during onboarding', 'Requested lazily after login'],
])

H('Part A — Practice, session player, results, history and my stats')
MD(A2+'A_practice_session_results_history_stats.md')
H('Part B — Onboarding, diagnostic, learning path, cheat sheets, case studies, flashcards')
MD(A2+'B_onboarding_diagnostic_path_cheatsheets_cases_flashcards.md')
H('Part C — Ask Cleo (AI tutor) and the Library')
MD(A2+'C_ask_cleo_and_library.md')
H('Part D — How pre-built material is made and delivered; monetisation; analytics')
MD(A2+'D_content_pipeline_monetisation_analytics.md')
H('Part E — Navigation, every button, settings, social and battles')
MD(A2+'E_navigation_buttons_settings_social_battles.md')
H('Part F — UI and design system')
MD(A2+'F_ui_and_design_system.md')
H('Part G — Client-side functions and algorithms')
MD(A2+'G_client_functions_and_algorithms.md')

H('6. Reference code: client kit (TypeScript, tested)')
P('An original implementation of the app\'s client-side rules, written from the observed thresholds and behaviour, not from decompiled code. Every module has unit tests (16 tests pass with node --test). It has not been run inside an app.')
T(['Module', 'What it does', 'Basis'], [
 ['keys.ts', 'Random and stable (replay-safe) idempotency keys; stable keys use four seeded 32-bit hashes joined with a unit separator', 'Observed behaviour; hash is ours'],
 ['http.ts', 'API client: base-URL failover on network errors and 502/503/504, 12 s timeout, one-shot token refresh shared by concurrent 401s, GET coalescing, backend error messages', 'Observed behaviour and constants'],
 ['offline.ts', 'Backoff (500 ms doubling, cap 8 s), outbox outcomes (2xx sent; 408/429/5xx/network retry; other 4xx discard), 24 h prefetch interval, 24 h offline window, chunk validation', 'Observed constants and rules'],
 ['session.ts', 'Remaining time with pause, battle clock skew, 50 s per question estimate, m:ss formatting', 'Observed'],
 ['labels.ts', 'Readiness (85/75/60), case readiness (80/60), difficulty (40/60/80), diagnostic buckets (40/60/80), Bloom to recall/apply/judgment, evidence minimums, streak milestones 3,7,14,30,50,100,200,365', 'Observed thresholds; level names are ours'],
 ['rank.ts', 'Nine-tier ladder with stars, tier-start rule, tier and stars from points, win/loss rule, combo multiplier', 'Tier names and star counts observed; the point rules are a design (the server decides in the app)'],
 ['flashcards.ts', 'Queue advance (good/easy leave, again/hard return), mastery percent, Fisher-Yates shuffle', 'Observed behaviour'],
 ['manila.ts', 'Fixed UTC+8 day clock, ms until next day, days until exam', 'Observed'],
])
NOTE('Parity: constants such as the combo cap and the middle forecast thresholds were not recovered, so those values are placeholders. Use this kit as a starting point and tune it with your own data.')
FILES(['product/client-kit/src/keys.ts','product/client-kit/src/http.ts','product/client-kit/src/offline.ts','product/client-kit/src/session.ts','product/client-kit/src/labels.ts','product/client-kit/src/rank.ts','product/client-kit/src/flashcards.ts','product/client-kit/src/manila.ts'])

H('7. What to take for the Nursing Study App')
T(['Copy the idea', 'Why'], [
 ['Base-URL failover, single-flight token refresh, request coalescing', 'Reliability on poor mobile networks; implemented and tested in http.ts'],
 ['Idempotency keys on every write', 'Safe retries from the offline outbox; keys.ts'],
 ['Offline bundle plus outbox with retry/discard rules', 'Study on hospital wards or commutes without signal; offline.ts'],
 ['Diagnostic that seeds the plan and readiness', 'Fast personalised start; labels.ts buckets'],
 ['Study plan with daily tasks that launch sessions, plus rebalance', 'Keeps learners on track to the exam date'],
 ['Blueprint taxonomy on every question and report; Bloom buckets (recall, apply, judgment)', 'Actionable feedback by exam area and thinking level'],
 ['Soft time limits and a tag-to-mastery action in the player', 'Lower anxiety; learner-driven review queues'],
 ['Fixed-timezone day clock for streaks and goals', 'Fair streaks; manila.ts'],
 ['Cleo-style assistant that curates from an approved bank first', 'Avoids unreviewed AI-written clinical items'],
])
P('Avoid or handle with care:', bold=True)
B([
 'Shipping answer keys and rationales to the device for offline use makes them extractable; decide deliberately (the app downloads them for premium users).',
 'Embedding personal data such as passer rosters in the app bundle.',
 'Session replay and ad SDKs in a student app without masking and consent.',
 'Heavy gamification and a long queue of launch prompts in a high-stakes exam context; offer a quiet focus mode.',
 'Developer screens and tooling left in release builds.',
 'Any AI-generated clinical content without expert review and a visible disclaimer.',
])

H('8. Limits of this research')
B([
 'All scoring, adaptive selection, mastery, passing-probability, ranking, planner rebalancing, AI prompts and the content pipeline are server-side and not visible. Statements about them are inferences.',
 'The app was never run: UI descriptions come from code and strings, not screenshots.',
 'The upload had no native libraries (lib/ folder), so LiveKit, WebRTC and Hermes native code were not inspected.',
 'Unresolved constants are listed at the end of each part (for example mock exam length, adaptive slider default, combo cap, some forecast thresholds).',
 'The sources are untrusted data and no instructions in them were followed.',
])
d.save(R+'Tangerine_Prep_Complete_Guide.docx')
print('saved')
