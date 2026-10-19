import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

A2 = '/home/user/readthischatgpt/Nursing app master plan/Content Authoring Kit/research/'
R = '/home/user/readthischatgpt/Nursing app master plan/Content Authoring Kit/'

d = Document()
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.9)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); tcPr.append(sh)

def H(t, l=1): d.add_heading(t, level=l)
def P(t, bold=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic; return p
def B(items):
    for t in items:
        p = d.add_paragraph(style='List Bullet')
        if isinstance(t, tuple):
            r = p.add_run(t[0]); r.bold = True; p.add_run(t[1])
        else: p.add_run(t)
def T(head, rows, widths=None):
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); shade(c, 'D9E2F3')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    d.add_paragraph()
def CODE(t):
    p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.15)
    r = p.add_run(t); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    pPr = p._p.get_or_add_pPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'F2F2F2'); pPr.append(sh)
def NOTE(t):
    p = d.add_paragraph(); r = p.add_run('Note: '); r.bold = True; p.add_run(t)


def inline(p, text):
    # **bold** and `code` runs
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else: p.add_run(part)

def MD(path, level_shift=1):
    lines = open(path, errors='replace').read().split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln or ln.startswith('Labels:') or ln.startswith('Source is untrusted'): i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            lvl = min(len(m.group(1)) + level_shift, 4)
            if len(m.group(1)) == 1: i += 1; continue   # file title replaced by our own section heading
            d.add_heading(m.group(2), level=lvl); i += 1; continue
        if ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells): rows.append(cells)
                i += 1
            if rows:
                n = max(len(r) for r in rows)
                t = d.add_table(rows=0, cols=n); t.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(n):
                        cells[ci].text = ''
                        p = cells[ci].paragraphs[0]
                        inline(p, r[ci] if ci < len(r) else '')
                        for run in p.runs: run.font.size = Pt(9); run.bold = run.bold or ri == 0
                        if ri == 0: shade(cells[ci], 'D9E2F3')
                d.add_paragraph()
            continue
        mb = re.match(r'^(\s*)[-*]\s+(.*)', ln)
        if mb:
            p = d.add_paragraph(style='List Bullet'); inline(p, mb.group(2)); i += 1; continue
        p = d.add_paragraph(); inline(p, ln.strip()); i += 1


def FILE(path, maxlines=None):
    t = open(path, encoding='utf-8').read().rstrip('\n')
    CODE(t)



def FILES(paths):
    for p in paths:
        P(p, bold=True); FILE(R+p)

d.add_heading('Content Authoring and Prompt Reference', 0)
P('How topics, lessons and questions are represented, authored and generated in the reference apps, and a working kit for the Nursing Study App.', italic=True)
P('Covers Tangerine Prep 1.9.0 (code fully decompiled), Gizmo 3.2.130 (string table only) and AnkiDroid 2.25.1 (assets and strings), plus the project\'s canonical PNLE Question Maker prompt (v2.0). Static analysis only; no app was run. Short quotes only; no proprietary prompt libraries or content were copied.')

H('1. Summary')
B([
 ('Did we read how questions and topics are made? ', 'We read how they are represented and requested. How they are authored is mostly server-side and not in any APK: no question-writing prompts, models or distractor logic are visible, except two small client-side prompts in Tangerine and a few explainer snippets in Gizmo.'),
 ('Correction to earlier wording. ', 'Tangerine\'s exam assistant does not write new exam questions. It says it curates existing questions from its question bank. Its library feature can generate flashcards and questions from a user\'s own sources, but those prompts run on the server and were not visible.'),
 ('Pre-built content is mostly fetched, not embedded. ', 'In Tangerine only the five-part taxonomy, a rank ladder, 24 loading tips, paywall and streak copy, and a few samples (6 demo MCQs, 1 sample report, 2 sample case studies) ship in the bundle. Cheat sheets, case studies, flashcards, topics and onboarding content come from the server.'),
 ('Your canonical prompt is already strong. ', 'It specifies exactly four options, one best answer, a rationale standard, distractor engineering and metadata. The kit turns it into a data contract (schemas and a validator) and adds six companion prompts for jobs it does not cover.'),
 ('Delivered and tested. ', 'JSON Schemas, a topic-tree builder (your hierarchy: 7 subjects, 636 nested topics, depth 7), a content validator, six sample items and six companion prompts. 11 of 11 tests pass. Samples are original and unverified (marked needs_review).'),
])

H('2. How topics are represented')
T(['Source', 'Structure', 'Where it lives'], [
 ['Tangerine Prep', 'Five Nursing Practices (NP1-NP5), each with a subject name, description, colour hue and icon > practice area (code, name, short name, icon) > topic (id, name) > TOS section (string) > tags. Separate axes: the practice area and the NP can differ for one item. Bloom level (remembering to creating) and difficulty are item attributes.', 'NP constants embedded; practice areas, topics and TOS sections fetched from /practice-areas/ and /tos/question-structure/ and cached in SQLite'],
 ['Gizmo', 'Decks and sub-decks, classes and school lists, and a very large list of exam and certification names used as presets. Social objects (groups, friends) wrap decks.', 'Mostly server-side; presets are strings in the bundle'],
 ['AnkiDroid / Anki', 'Decks nested with "Parent::Child", hierarchical tags, filtered decks for custom study, note types for card layouts. Topic = deck and/or tag.', 'Collection database on the device'],
 ['Your repo', '7 subjects (Community Health, Fundamentals, Medical-Surgical, OB, PALMER, Pediatric, Psychiatric) and 636 nested topic folders, depth up to 7, from the PNLE TOP 1 Anki deck.', 'nursing preset topics/Anki Deck Topic Hierarchy; machine-readable copy: samples/topic-tree.generated.json'],
])
NOTE('The folder README counts 637 nested topics; counting directories on disk gives 636 (644 directories including the root and the 7 subjects). The 7 subjects do not map one-to-one onto NP1-NP5; the kit includes a provisional mapping that must be checked against the official TOS.')

H('3. How questions are represented')
H('3.1 Tangerine Prep (observed in code)', 2)
B([
 ('Standard practice question. ', 'One stem field named body, an options array of {id, option_text, option_order, is_correct}, a topic, a practice area, tags, bloom_level, difficulty, is_high_yield and is_verified flags and a verified reviewer badge. A separate answer block carries the correct option id, rationale, a simplified rationale, sources, images (with credit and licence), optional audio, and cohort statistics (mean accuracy and time). The player renders single-select options only.'),
 ('Case-study (situation) items. ', 'A shared stem plus items of type mcq, sata, matrix, highlight, cloze and free_response. Grading returns earned and maximum units and credits, so partial credit exists. No ordering, bowtie or exhibit types were found.'),
 ('Learner tools on each question. ', 'Elimination mode, rationale in full or simplified form, a report-this-question action (wording confusing, too long, incorrect, other), per-question mastery and an offline cache with a content hash.'),
])
CODE("""// Standard question (observed shape, simplified)
{ "question": { "id": 456, "body": "stem...", "options": [{ "id": 1, "option_text": "...", "option_order": 1, "is_correct": false }],
    "topic": {"name": "..."}, "practice_area": {"short_name": "..."}, "tags": [{"id": 9, "name": "..."}],
    "bloom_level": "applying", "difficulty": "medium", "is_verified": true,
    "answer": { "correct_option_id": 1, "rationale": "...", "simplified_rationale": "...",
                "answer_sources": [], "rationale_images": [], "mean_accuracy": 0.0 } } }""")
H('3.2 Gizmo (observed names, assembled shape inferred)', 2)
B([
 'Every question is an Activity whose public content (lead-in, option banks, items, positions, blanks) is separate from a redacted answer key. A question has Segments (parts that can be answered) and Inputs (an answering method: choice, typing, ordering, matching, cloze or flashcard) with a commit policy, plus evaluation, mistake, skip, heart and hint policies owned by the activity.',
 'Text markup marks answerable parts with [text](#segment-id), plus latex and chart tags. Free-response marking uses student-answer, mark-scheme and marks-awarded tags.',
 'A few embedded prompt snippets exist: a wrong-step misconception explainer and a next-step explainer for procedural questions, and chart-planning instructions. There is no tutor persona text in the bundle.',
])
H('3.3 AnkiDroid / Anki', 2)
B([
 'A note has a note type (fields plus card templates). Stock types include Basic, Basic (and reversed card), Basic (optional reversed card), Basic (type in the answer), Cloze and Image Occlusion.',
 'Cloze syntax is {{c1::text::hint}}; one card is generated per cloze number. Template syntax includes {{Field}}, {{FrontSide}}, {{type:Field}}, {{cloze:Field}}, and {{#Field}}...{{/Field}} conditionals.',
 'Multiple choice is not a built-in type. The usual community convention is a note type with fields for the question, the options, the answer and an explanation, with a little JavaScript in the template. The AnkiDroid card scripting API (js-api.js v0.0.3) lets a template call the app.',
 'Text import supports separator and HTML directives; deck and tag columns could not be confirmed in the files.',
])
H('3.4 Side by side', 2)
T(['Aspect', 'Tangerine', 'Gizmo', 'Anki', 'Canonical PNLE prompt'], [
 ['Standard MCQ', '4+ options, single select', 'Choice input', 'Community convention', 'Exactly 4 (A-D), one best answer'],
 ['Multi-select', 'SATA in case studies', 'Multiple-choice input', 'Not built in', 'Not covered'],
 ['Matrix / cloze / highlight', 'Yes (case studies)', 'Cloze, matching, ordering', 'Cloze built in', 'Not covered'],
 ['Rationale', 'Full + simplified + audio + images + sources', 'Explain flow', 'Free text field', 'Rationale standard (reason, cue, distractor, citation)'],
 ['Hints', 'Elimination mode', 'Hint policy', 'Template hint filter', 'Not covered'],
 ['Metadata', 'Topic, area, tags, Bloom, difficulty, verified', 'Deck/class/topic', 'Deck, tags', 'Subject, TOS area, topic, competency, cognitive domain, difficulty, process phase, evidence status'],
 ['Quality signals', 'Verified badge, report item, cohort stats', 'Redacted answer key', 'None', 'QC loop, source rules'],
])

H('4. How questions are generated in the apps')
T(['App', 'Feature', 'What the client shows', 'Where the intelligence is'], [
 ['Tangerine', 'Ask Cleo exam builder', 'The server returns an exam_draft with a generation_request; the client posts it unchanged to /sessions/ai-generate/v2/, then creates the session. The assistant states it curates from the existing question bank.', 'Server (selection and any LLM call)'],
 ['Tangerine', 'AI Curator and daily sets', 'Optional focus text, then /sessions/ai-generate/ ; daily kinds daily_climb and daily_polish', 'Server'],
 ['Tangerine', 'Adaptive exam', 'Session config carries model name, model version, repetition mode, allocated count and a fallback reason such as learner_cold_start', 'Server learner model'],
 ['Tangerine', 'Library generation', 'POST /library/folders/{id}/generate/ with kind (flashcard, question, Explain It, cheat sheet), count, focus and source ids; refine with an instruction', 'Server prompts (not visible)'],
 ['Tangerine', 'Client-side prompts', 'Two prompts built in the app: a stats-analysis prompt that asks for strict JSON, and a one-sentence acknowledgement prompt. User text is interpolated into the prompt, a prompt-injection surface.', 'Server LLM via /sessions/cleo-reply/'],
 ['Gizmo', 'Procedural questions, study notes, magic import, daily quiz, explainers', 'Endpoints and schema names only; field lists are candidates', 'ai.gizmo.ai worker'],
 ['AnkiDroid', 'None', 'No generation; content comes from users and imported decks', 'n/a'],
])
H('4.1 Pre-built prompts and chips found in Tangerine (summary)', 2)
B([
 'Suggested chips on home: create an exam, analyse my performance, build an exam from my blind spots, what should I study today, create pharmacology flashcards, explain a concept simply.',
 'Per-question chips: explain simply, create a flashcard, why are the other options wrong, quiz me on this concept, how would this appear clinically, what should I review next.',
 'Cleo action types returned by the server: performance_visualization, ask_user_input, exam_draft, support_ticket, cheat_sheet_recommendations, create_flashcards. Languages offered: English, Filipino, Bisaya.',
 'Pattern worth copying: the exam builder first asks what to base the exam on (weak spots, strong skills, or a chosen topic) before generating.',
])

H('5. Pre-built topics and lessons in Tangerine')
T(['Category', 'Embedded or fetched', 'Notes'], [
 ['Five Nursing Practices', 'Embedded', 'NP1 Community Health; NP2 Maternal, Child & Human Development; NP3 Surgery, Oxygenation & Fluid Balance; NP4 Nutrition, Endocrine & Coordination; NP5 Mental Health & Acute Care; each with a one-sentence description, hue and icon'],
 ['Practice areas and topics', 'Fetched and cached', 'Practice-area and topic tables in the offline SQLite database'],
 ['TOS sections', 'Fetched', 'Only one section name is embedded in the bundle'],
 ['Cheat sheets', 'Fetched (catalog plus offline cache)', 'Paged, zoomable documents; a four-card intro is embedded; premium feature; titles and counts are not in the bundle'],
 ['Case studies / situations', 'Fetched; 2 samples embedded', 'Each sample has 5 items'],
 ['Flashcards', 'Fetched', 'No embedded decks'],
 ['Demo questions', 'Embedded', '6 demo MCQs and 1 sample exam report'],
 ['Rank ladder', 'Embedded', '9 tiers from Student II to Apex Topnotcher'],
 ['Loading tips, streak copy, paywall copy', 'Embedded', '24 loading tips'],
 ['Notifications', 'Server-pushed', 'Only type keys and settings labels are embedded'],
 ['Onboarding configuration', 'Fetched', '/app-platform/onboarding-config/'],
])
P('There are no embedded lesson scripts or notification templates in the bundle. Question counts, cheat-sheet titles and the topic list therefore cannot be measured from the APK. The exam-passer name list is described elsewhere and deliberately not reproduced.')

H('6. Gap analysis: your canonical prompt versus the reference apps')
P('Strengths of your prompt (none of the apps show equivalent rules): official-framework check before generation, evidence rules, a pre-generation blueprint, distractor engineering, a rationale standard, Philippine practice context, clinical prioritisation, a mandatory QC loop and strict output modes.')
P('Things the apps do that your prompt does not yet cover:', bold=True)
T(['Gap', 'Seen in', 'Suggested handling in the kit'], [
 ['Machine-readable output', 'All apps use one structured item object', 'schemas/question.schema.json and prompts/01'],
 ['Multi-select, matrix, cloze, highlight items', 'Tangerine case studies', 'Schema extensions and prompts/02; the practice-engine scores them'],
 ['Second-pass review and learner reporting', 'Verified badge, report-this-question', 'prompts/03; a report action in the app'],
 ['Simplified rationale, hint, audio script', 'Tangerine rationale modes and audio; Gizmo hints', 'prompts/04'],
 ['Curating sets from an approved bank', 'Tangerine "AI" exam builder', 'prompts/05 or the Rasch selector in App 4 product/practice-engine'],
 ['Flashcards that import into Anki', 'Tangerine create_flashcards; Anki note types', 'prompts/06 plus App 2 product/apkg-builder'],
 ['Image credit and licence fields', 'Tangerine rationale images', 'images[] in the schema (url, alt_text, caption, credit, license)'],
 ['Topic ids', 'Tangerine topic_id / tos_section', 'topic_id validated against topic-tree.json'],
])

H('7. The kit')
T(['Component', 'What it does', 'Verified by'], [
 ['schemas/topic-tree.schema.json', 'Subjects, topics, ids as slugs, optional provisional NP mapping', 'Built tree validates; ids unique'],
 ['schemas/question.schema.json', 'MCQ (exactly 4 options A-D, one answer) plus sata, matrix, cloze, highlight and situation clusters; metadata fields from the canonical prompt', 'Samples validate; bad items rejected'],
 ['tools/build_topic_tree.py', 'Converts your folder hierarchy to topic-tree.json', 'Test: node count equals directories on disk (643 = 7 + 636), depth 7; committed JSON is current'],
 ['tools/validate_content.py', 'Schema plus semantic rules: option labels and uniqueness, banned all/none of the above forms, option length ratio, absolute words, missing distractor notes, index ranges, cloze blank consistency, highlight rules, duplicate stems, answer-key skew, longest-answer bias, unknown topic ids, publish gate for needs_review', '11 tests pass'],
 ['samples/', '6 original items (MCQ, SATA, cloze, a diabetes situation with an MCQ and a matrix, and a pressure-injury MCQ), all needs_review', 'Validate with 0 errors and 6 review warnings'],
 ['prompts/ (6 files)', 'Companion prompts under the canonical prompt', 'Reviewed by reading only; not run against a model'],
])
P('Pipeline', bold=True)
B([
 '1. Choose a topic id from topic-tree.json.',
 '2. Generate with the canonical prompt plus prompts/01 (JSON output).',
 '3. Run validate_content.py with the topic tree; fix errors.',
 '4. Run prompts/03 as a second-pass reviewer.',
 '5. A clinician or instructor reviews and sets evidence_status; only then publish (the --publish flag refuses needs_review).',
])
NOTE('The sample items are original and unverified. They test the pipeline and are not study material. The subject-to-NP mapping is provisional.')

H('8. What could not be determined')
B([
 'How the apps actually write or select questions: prompts, models, distractor logic and selection formulas are server-side.',
 'The full generation_request field list in Tangerine (the server authors it) and the request fields of Gizmo\'s AI endpoints (only candidate names in the string table).',
 'Counts of questions, cheat sheets and lessons in any app, because they are fetched.',
 'Anki import header directives beyond separator and HTML, and the full qfmt/afmt of stock note types (inferred from upstream).',
 'Whether the kit\'s companion prompts produce good items: they have not been run against a model.',
])

H('9. Addendum A — Tangerine: questions, topics, AI requests, prompts')
MD(A2+'tangerine_questions_topics_prompts.md')
H('10. Addendum B — Tangerine: pre-built and embedded content')
MD(A2+'tangerine_prebuilt_content.md')
H('11. Addendum C — Gizmo: questions and prompts')
MD(A2+'gizmo_questions_prompts.md')
H('12. Addendum D — AnkiDroid: note types and authoring')
MD(A2+'ankidroid_note_types_authoring.md')

d.add_page_break()
H('Appendix — Kit source (schemas, tools and prompts)')
FILES(['schemas/topic-tree.schema.json','schemas/question.schema.json','tools/build_topic_tree.py','tools/validate_content.py','samples/questions.sample.json'])
for f in ['01_machine_readable_output','02_case_study_items','03_second_pass_reviewer','04_rationale_variants_and_hints','05_exam_set_curation','06_flashcards_from_notes']:
    FILES(['prompts/'+f+'.md'])
d.save(R+'Content_Authoring_and_Prompt_Reference.docx')
print('saved')
