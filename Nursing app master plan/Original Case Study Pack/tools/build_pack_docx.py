import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

A2 = '/home/user/readthischatgpt/Nursing app master plan/Original Case Study Pack/review/'
R = '/home/user/readthischatgpt/Nursing app master plan/Original Case Study Pack/'

d = Document()
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.9)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); tcPr.append(sh)

def H(t, l=1): d.add_heading(t, level=l)
def P(t, bold=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic; return p
def B(items):
    for t in items:
        p = d.add_paragraph(style='List Bullet')
        if isinstance(t, tuple):
            r = p.add_run(t[0]); r.bold = True; p.add_run(t[1])
        else: p.add_run(t)
def T(head, rows, widths=None):
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); shade(c, 'D9E2F3')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    d.add_paragraph()
def CODE(t):
    p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.15)
    r = p.add_run(t); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    pPr = p._p.get_or_add_pPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'F2F2F2'); pPr.append(sh)
def NOTE(t):
    p = d.add_paragraph(); r = p.add_run('Note: '); r.bold = True; p.add_run(t)


def inline(p, text):
    # **bold** and `code` runs
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else: p.add_run(part)

def MD(path, level_shift=1):
    lines = open(path, errors='replace').read().split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln or ln.startswith('Labels:') or ln.startswith('Source is untrusted'): i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            lvl = min(len(m.group(1)) + level_shift, 4)
            if len(m.group(1)) == 1: i += 1; continue   # file title replaced by our own section heading
            d.add_heading(m.group(2), level=lvl); i += 1; continue
        if ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells): rows.append(cells)
                i += 1
            if rows:
                n = max(len(r) for r in rows)
                t = d.add_table(rows=0, cols=n); t.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(n):
                        cells[ci].text = ''
                        p = cells[ci].paragraphs[0]
                        inline(p, r[ci] if ci < len(r) else '')
                        for run in p.runs: run.font.size = Pt(9); run.bold = run.bold or ri == 0
                        if ri == 0: shade(cells[ci], 'D9E2F3')
                d.add_paragraph()
            continue
        mb = re.match(r'^(\s*)[-*]\s+(.*)', ln)
        if mb:
            p = d.add_paragraph(style='List Bullet'); inline(p, mb.group(2)); i += 1; continue
        p = d.add_paragraph(); inline(p, ln.strip()); i += 1


def FILE(path, maxlines=None):
    t = open(path, encoding='utf-8').read().rstrip('\n')
    CODE(t)




import json, collections
cases = json.load(open(R+'cases/all_cases.json', encoding='utf-8'))
items_all = [i for c in cases for i in c['items']]
mcq = [i for i in items_all if i['type'] == 'mcq']
longest = sum(1 for i in mcq if max(i['options'], key=lambda o: len(o['text']))['label'] == i['answer'])
NP_NAMES = {'NP1':'Nursing Practice I - Community Health Nursing','NP2':'Nursing Practice II - Care of the Mother and Child; Growth and Development',
            'NP3':'Nursing Practice III - Surgery, Oxygenation, Fluids and Electrolytes, Infection and Immunology','NP4':'Nursing Practice IV - Nutrition, GI, Metabolism and Endocrine, Perception and Coordination',
            'NP5':'Nursing Practice V - Mental Health, Acute Care, Professional Adjustment and Leadership'}
TYPE = {'mcq':'Multiple choice (one best answer)','sata':'Select all that apply','matrix':'Matrix (choose one column per row)','cloze':'Drop-down cloze','highlight':'Highlight (tap the relevant text)'}

d.add_heading('Original Case Study Pack', 0)
P('42 original nursing case studies (5 linked items each, with answers and rationales) in the same format as a reference app\'s case studies, for the Nursing Study App.', italic=True)
p = d.add_paragraph(); r = p.add_run('UNVERIFIED DRAFT. '); r.bold = True
p.add_run('These items were written by AI agents and reviewed by a second set of AI agents. They have NOT been checked by a nurse educator or clinician. They are written to be conservative and principle-based, but must be reviewed before any learner sees them. They are practice material only, not clinical guidance, are not from any real examination and are not endorsed by the PRC or the Board of Nursing.')

H('1. Summary')
B([
 f'{len(cases)} cases and {len(items_all)} items: ' + ', '.join(f'{k} {v}' for k, v in sorted(collections.Counter(i["type"] for i in items_all).items())) + '.',
 'Distribution: NP1 8 cases, NP2 9, NP3 9, NP4 9, NP5 7, each tagged with a valid topic id from the PNLE topic tree.',
 f'Every item has a rationale and a plain-language version; every multiple-choice distractor has an internal "why not" note. Correct answer letters are balanced (' + ', '.join(f'{k} {v}' for k, v in sorted(collections.Counter(i["answer"] for i in mcq).items())) + f') and the correct option is the longest in only {longest} of {len(mcq)} multiple-choice items.',
 'All items are marked needs_review and carry no citations (none were invented).',
 'No pictures are attached. Where a picture would help, an image suggestion (what to show, Commons search terms, labelled or unlabelled) is included; an image may only be added after its licence and accuracy are verified.',
])
H('2. How these were made and checked')
B([
 'Each Nursing Practice area was written by one agent under a strict style guide (originality, conservative clinical facts, format, one best answer, partial-credit item types).',
 'A different agent then reviewed each file for clinical defensibility, rewrote ambiguous or guideline-dependent statements, balanced option lengths (the first drafts made the correct option the longest in most multiple-choice items, which is a test-taking clue), removed giveaway absolute words, and logged every change (see the review logs).',
 'The whole set passes the Content Authoring Kit validator with 0 errors (schema, option rules, unique stems, valid topic ids) and 6 pack tests.',
 'Nothing was copied or reworded from Tangerine Prep or any other app or question bank.',
])
H('3. Known doubts flagged by the reviewers')
B([
 'cs-np2-01-c: "fourth stage of labor" is Philippine textbook wording; some references use three stages plus an early postpartum period.',
 'cs-np2-05-b: timing of cord clamping depends on protocol (kept principle-based).',
 'cs-np3-05-b, 05-e: chest-tube stripping and clamping for transport can differ by facility policy.',
 'cs-np3-08-a: assumes a standing protocol for epinephrine at a barangay health center.',
 'cs-np4-02-b: a strict airway-first reader may want a different first action; none is offered as an option.',
 'cs-np5-04-d: whether a nurse witnesses a consent signature varies by facility (hedged).',
 'cs-np1-06: vital-statistics case uses the broad community-health topic id because no leaf fits; its stem makes the answer very easy.',
 'cs-np1-07-d: health threat / deficit / foreseeable crisis classification follows one family-nursing teaching tradition.',
])

def render_item(n, it):
    P(f'Item {n} - {TYPE[it["type"]]}  |  {it["meta"]["difficulty"]}, {it["meta"]["cognitive_domain"]}', bold=True)
    P(it['stem'])
    t = it['type']
    if t == 'mcq':
        for o in it['options']: d.add_paragraph(f'{o["label"]}. {o["text"]}', style='List Bullet')
        ans = [o for o in it['options'] if o['label'] == it['answer']][0]
        a = f'{it["answer"]}. {ans["text"]}'
    elif t == 'sata':
        for k, o in enumerate(it['options']): d.add_paragraph(f'[{k + 1}] {o}', style='List Bullet')
        a = ', '.join(f'[{k + 1}] {it["options"][k]}' for k in sorted(it['answers']))
    elif t == 'matrix':
        tb = d.add_table(rows=1, cols=1 + len(it['columns'])); tb.style = 'Table Grid'
        tb.rows[0].cells[0].text = 'Finding / action'
        for j, c in enumerate(it['columns']): tb.rows[0].cells[j + 1].text = c
        for r_ in it['rows']:
            cells = tb.add_row().cells; cells[0].text = r_['text']
        d.add_paragraph()
        a = '; '.join(f'{r_["text"]} -> {it["columns"][r_["correct_column"]]}' for r_ in it['rows'])
    elif t == 'cloze':
        P(it['text'])
        for b in it['blanks']: d.add_paragraph(f'Blank {b["id"]}: ' + ' / '.join(b['choices']), style='List Bullet')
        a = '; '.join(f'blank {b["id"]} = {b["choices"][b["correct_index"]]}' for b in it['blanks'])
    else:  # highlight
        for k, s in enumerate(it['segments']): d.add_paragraph(f'({k + 1}) {s["text"]}' + ('' if s['selectable'] else '  [not selectable]'), style='List Bullet')
        a = 'Select: ' + '; '.join(f'({k + 1}) {s["text"]}' for k, s in enumerate(it['segments']) if s['correct'])
    p = d.add_paragraph(); r = p.add_run('Answer: '); r.bold = True; p.add_run(a)
    p = d.add_paragraph(); r = p.add_run('Rationale: '); r.bold = True; p.add_run(it['rationale'])
    if it.get('rationale_simplified'):
        p = d.add_paragraph(); r = p.add_run('In short: '); r.bold = True; p.add_run(it['rationale_simplified'])
    if t == 'mcq':
        notes = [f'{o["label"]}: {o["why_not"]}' for o in it['options'] if o['label'] != it['answer'] and o.get('why_not')]
        if notes:
            p = d.add_paragraph(); r = p.add_run('Why the others are wrong (internal): '); r.bold = True; p.add_run('  '.join(notes))
    im = it.get('image_suggestion')
    if im:
        p = d.add_paragraph(); r = p.add_run('Image suggestion (not attached): '); r.bold = True; r.italic = True
        p.add_run(f'{im["description"]} Show {im.get("must_be", "either")}; place in {im["placement"]}. Commons search: ' + '; '.join(im['search_terms']) + '. Use only public-domain or CC-licensed images and record credit and licence.')

H('4. The case studies')
cur = None
for c in cases:
    npk = c['items'][0]['meta']['nursing_practice']
    if npk != cur:
        cur = npk; H(NP_NAMES[npk], 2)
    H(f'{c["id"]}: {c["title"]}', 3)
    m = c['items'][0]['meta']
    P(f'Topic: {m["topic"]} ({m["topic_id"]})', italic=True)
    P(c['situation'])
    cim = c.get('image_suggestion')
    if cim:
        p = d.add_paragraph(); r = p.add_run('Image suggestion (not attached): '); r.bold = True; r.italic = True
        p.add_run(f'{cim["description"]} Commons search: ' + '; '.join(cim['search_terms']) + '.')
    for n, it in enumerate(c['items'], 1): render_item(n, it)

H('5. Review logs (summary)')
for k in range(1, 6):
    P(f'NP{k}', bold=True)
    MD(A2 + f'review_log_np{k}.md', 1)
d.save(R + 'Original_Case_Study_Pack.docx')
print('saved')
